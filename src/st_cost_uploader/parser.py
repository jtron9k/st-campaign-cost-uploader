"""Spreadsheet reading and layout normalization.

Two layouts are accepted:
  long: one row per campaign-month, with campaign / period / amount columns
  wide: campaigns down the left, one column per month

Both normalize to SheetRow. Detection is reported to the caller so the UI can
show it and offer an override, because a misdetected layout must never be
silent.
"""

from __future__ import annotations

import calendar
import csv
import io
import re
from datetime import date

from openpyxl import load_workbook

from st_cost_uploader.models import ParseError, ParseResult, SheetRow
from st_cost_uploader.money import parse_money

MIN_PERIOD_COLUMNS_FOR_WIDE = 2

_MONTHS: dict[str, int] = {}
for _i in range(1, 13):
    _MONTHS[calendar.month_name[_i].lower()] = _i
    _MONTHS[calendar.month_abbr[_i].lower()] = _i

_CAMPAIGN_PATTERN = re.compile(r"campaign|name", re.IGNORECASE)
_PERIOD_PATTERN = re.compile(r"month|period|date", re.IGNORECASE)
_AMOUNT_PATTERN = re.compile(r"spend|cost|total|amount|budget", re.IGNORECASE)


class ParserError(Exception):
    """The file cannot be read at all."""


def parse_period(value: object) -> tuple[int, int] | None:
    """Coerce a cell into (year, month). Returns None when it is not a period."""
    if value is None:
        return None
    if isinstance(value, date):
        return value.year, value.month

    text = str(value).strip()
    if not text:
        return None

    m = re.fullmatch(r"(\d{4})[-/](\d{1,2})(?:[-/]\d{1,2})?", text)
    if m:
        year, month = int(m.group(1)), int(m.group(2))
        return (year, month) if 1 <= month <= 12 else None

    m = re.fullmatch(r"(\d{1,2})[-/](\d{4})", text)
    if m:
        month, year = int(m.group(1)), int(m.group(2))
        return (year, month) if 1 <= month <= 12 else None

    m = re.fullmatch(r"([A-Za-z]+)[\s,\-]+(\d{4})", text)
    if m:
        month = _MONTHS.get(m.group(1).lower())
        if month:
            return int(m.group(2)), month

    return None


def detect_layout(headers: list[object]) -> str:
    """Wide when two or more headers parse as periods, otherwise long."""
    period_columns = sum(1 for h in headers[1:] if parse_period(h) is not None)
    return "wide" if period_columns >= MIN_PERIOD_COLUMNS_FOR_WIDE else "long"


def _read_rows(data: bytes, filename: str) -> list[list[object]]:
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig")
        return [list(r) for r in csv.reader(io.StringIO(text))]

    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheet = workbook.active
    return [list(r) for r in sheet.iter_rows(values_only=True)]


def _find_column(headers: list[object], pattern: re.Pattern[str]) -> int | None:
    for index, header in enumerate(headers):
        if header is not None and pattern.search(str(header)):
            return index
    return None


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _parse_long(
    headers: list[object], body: list[tuple[int, list[object]]]
) -> tuple[list[SheetRow], list[ParseError], dict[str, str]]:
    campaign_col = _find_column(headers, _CAMPAIGN_PATTERN)
    period_col = _find_column(headers, _PERIOD_PATTERN)
    amount_col = _find_column(headers, _AMOUNT_PATTERN)

    missing = [
        label
        for label, col in (
            ("campaign", campaign_col), ("period", period_col), ("amount", amount_col)
        )
        if col is None
    ]
    if missing:
        raise ParserError(
            f"Long layout needs a {', '.join(missing)} column. "
            f"Found headers: {[_clean(h) for h in headers if _clean(h)]}"
        )

    mapping = {
        _clean(headers[campaign_col]): "campaign_name",
        _clean(headers[period_col]): "period",
        _clean(headers[amount_col]): "monthly_total",
    }

    rows: list[SheetRow] = []
    errors: list[ParseError] = []
    for line, raw in body:
        name = _clean(raw[campaign_col] if campaign_col < len(raw) else None)
        period = parse_period(raw[period_col] if period_col < len(raw) else None)
        amount = parse_money(raw[amount_col] if amount_col < len(raw) else None)

        if not name and period is None and amount is None:
            continue
        if not name:
            errors.append(ParseError(line, "Missing campaign name"))
            continue
        if period is None:
            errors.append(ParseError(line, f"Could not read a year and month for '{name}'"))
            continue
        if amount is None:
            errors.append(ParseError(line, f"Could not read an amount for '{name}'"))
            continue

        rows.append(SheetRow(name, period[0], period[1], amount, line))

    return rows, errors, mapping


def _parse_wide(
    headers: list[object], body: list[tuple[int, list[object]]]
) -> tuple[list[SheetRow], list[ParseError], dict[str, str]]:
    periods = [(i, p) for i, h in enumerate(headers) if (p := parse_period(h)) is not None]
    if not periods:
        raise ParserError(
            "Wide layout needs at least one column header that reads as a month, "
            f"such as 'Feb 2026'. Found: {[_clean(h) for h in headers if _clean(h)]}"
        )

    mapping = {_clean(headers[0]): "campaign_name"}
    for index, (year, month) in periods:
        mapping[_clean(headers[index])] = f"period:{year}-{month:02d}"

    rows: list[SheetRow] = []
    errors: list[ParseError] = []
    for line, raw in body:
        name = _clean(raw[0] if raw else None)
        if not name:
            if any(_clean(c) for c in raw):
                errors.append(ParseError(line, "Missing campaign name"))
            continue

        for index, (year, month) in periods:
            cell = raw[index] if index < len(raw) else None
            if _clean(cell) == "":
                continue
            amount = parse_money(cell)
            if amount is None:
                errors.append(
                    ParseError(line, f"Could not read an amount for '{name}' in {year}-{month:02d}")
                )
                continue
            rows.append(SheetRow(name, year, month, amount, line))

    return rows, errors, mapping


def parse_workbook(data: bytes, filename: str, layout: str | None = None) -> ParseResult:
    """Read a spreadsheet into canonical rows.

    layout: pass "long" or "wide" to override detection. None auto-detects.
    """
    raw_rows = _read_rows(data, filename)
    # Capture each row's real 1-based spreadsheet line before filtering out
    # blanks, so that reported source_row values always match the line the
    # marketer would see in their own spreadsheet, blank rows included.
    indexed_rows = [
        (index, row) for index, row in enumerate(raw_rows, start=1) if any(_clean(c) for c in row)
    ]
    if not indexed_rows:
        raise ParserError("The file has no rows.")

    (_, headers), body = indexed_rows[0], indexed_rows[1:]
    chosen = layout or detect_layout(headers)
    if chosen not in ("long", "wide"):
        raise ParserError(f"Unknown layout '{chosen}'. Use 'long' or 'wide'.")

    if chosen == "wide":
        rows, errors, mapping = _parse_wide(headers, body)
    else:
        rows, errors, mapping = _parse_long(headers, body)

    return ParseResult(
        layout=chosen,
        rows=tuple(rows),
        errors=tuple(errors),
        column_mapping=mapping,
    )
