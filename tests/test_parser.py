import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import Workbook

from st_cost_uploader.parser import detect_layout, parse_period, parse_workbook


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("2026-02", (2026, 2)),
        ("2026/2", (2026, 2)),
        ("2026-02-01", (2026, 2)),
        ("02-2026", (2026, 2)),
        ("Feb 2026", (2026, 2)),
        ("February 2026", (2026, 2)),
        ("feb 2026", (2026, 2)),
        ("Jan, 2026", (2026, 1)),
        (date(2026, 2, 14), (2026, 2)),
        ("2026-13", None),
        ("Campaign", None),
        ("", None),
        (None, None),
    ],
)
def test_parse_period(raw, expected):
    assert parse_period(raw) == expected


def test_detect_layout_long():
    assert detect_layout(["Campaign", "Month", "Spend"]) == "long"


def test_detect_layout_wide():
    assert detect_layout(["Campaign", "Jan 2026", "Feb 2026", "Mar 2026"]) == "wide"


def test_detect_layout_wide_needs_two_periods():
    # A single period-looking column is not enough to call it wide.
    assert detect_layout(["Campaign", "Feb 2026", "Notes"]) == "long"


def test_parse_long_layout():
    data = _xlsx([
        ["Campaign", "Month", "Spend"],
        ["Google - Search Brand", "2026-02", "5,000.00"],
        ["Yelp Ads", "2026-03", "$1,860.00"],
    ])
    result = parse_workbook(data, "spend.xlsx")

    assert result.layout == "long"
    assert result.errors == ()
    assert len(result.rows) == 2
    assert result.rows[0].campaign_name == "Google - Search Brand"
    assert result.rows[0].year == 2026
    assert result.rows[0].month == 2
    assert result.rows[0].monthly_total == Decimal("5000.00")
    assert result.rows[1].monthly_total == Decimal("1860.00")


def test_parse_wide_layout_unpivots():
    data = _xlsx([
        ["Campaign", "Jan 2026", "Feb 2026"],
        ["Yelp", "1000", "2000"],
        ["Nextdoor", "500", ""],
    ])
    result = parse_workbook(data, "spend.xlsx")

    assert result.layout == "wide"
    assert len(result.rows) == 3  # blank cell is skipped, not an error
    assert (result.rows[0].campaign_name, result.rows[0].month) == ("Yelp", 1)
    assert result.rows[1].monthly_total == Decimal("2000")
    assert (result.rows[2].campaign_name, result.rows[2].month) == ("Nextdoor", 1)


def test_manual_layout_override_beats_detection():
    # Headers look long, but the caller insists on wide.
    data = _xlsx([
        ["Campaign", "Jan 2026", "Notes"],
        ["Yelp", "1000", "ignore me"],
    ])
    result = parse_workbook(data, "spend.xlsx", layout="wide")

    assert result.layout == "wide"
    assert len(result.rows) == 1
    assert result.rows[0].month == 1


def test_bad_rows_become_errors_not_exceptions():
    data = _xlsx([
        ["Campaign", "Month", "Spend"],
        ["Good", "2026-02", "100"],
        ["", "2026-02", "100"],
        ["Bad Period", "not a date", "100"],
        ["Bad Amount", "2026-02", "n/a"],
    ])
    result = parse_workbook(data, "spend.xlsx")

    assert len(result.rows) == 1
    assert len(result.errors) == 3
    assert {e.source_row for e in result.errors} == {3, 4, 5}
    assert "campaign name" in result.errors[0].message.lower()


def test_column_mapping_is_reported():
    data = _xlsx([
        ["Campaign", "Month", "Spend"],
        ["Yelp", "2026-02", "100"],
    ])
    result = parse_workbook(data, "spend.xlsx")
    assert result.column_mapping == {
        "Campaign": "campaign_name",
        "Month": "period",
        "Spend": "monthly_total",
    }


def test_csv_is_supported():
    csv = b"Campaign,Month,Spend\nYelp,2026-02,1860.00\n"
    result = parse_workbook(csv, "spend.csv")
    assert result.layout == "long"
    assert result.rows[0].monthly_total == Decimal("1860.00")


def test_empty_file_raises():
    from st_cost_uploader.parser import ParserError

    with pytest.raises(ParserError):
        parse_workbook(_xlsx([]), "empty.xlsx")


def test_long_layout_missing_required_column_raises():
    from st_cost_uploader.parser import ParserError

    data = _xlsx([["Campaign", "Notes"], ["Yelp", "hi"]])
    with pytest.raises(ParserError):
        parse_workbook(data, "spend.xlsx")


def test_blank_rows_do_not_shift_reported_row_numbers():
    # Row 1: header. Row 2: good. Row 3: genuinely blank (dropped by the
    # blank-row filter). Row 4: bad period. A naive "offset + 2 against the
    # already-filtered body" numbering would report this as row 3 (it becomes
    # the second body row after filtering), which is the wrong line in the
    # user's actual spreadsheet. The correct report is row 4.
    data = _xlsx([
        ["Campaign", "Month", "Spend"],
        ["Good", "2026-02", "100"],
        ["", "", ""],
        ["Bad Period", "not a date", "100"],
    ])
    result = parse_workbook(data, "spend.xlsx")

    assert len(result.rows) == 1
    assert len(result.errors) == 1
    assert result.errors[0].source_row == 4


def test_blank_rows_do_not_shift_reported_row_numbers_wide():
    # Same idea in wide layout: row 3 is genuinely blank and gets filtered,
    # so the bad amount on row 4 must still be reported as row 4.
    data = _xlsx([
        ["Campaign", "Jan 2026", "Feb 2026"],
        ["Yelp", "1000", "2000"],
        ["", "", ""],
        ["Nextdoor", "n/a", "500"],
    ])
    result = parse_workbook(data, "spend.xlsx")

    assert len(result.errors) == 1
    assert result.errors[0].source_row == 4


def test_the_workbook_handle_is_closed_after_reading(monkeypatch):
    # load_workbook(read_only=True) holds an open zip archive that openpyxl
    # expects the caller to close. The rows must still be fully materialised
    # before that happens, because iter_rows is lazy in read-only mode.
    from openpyxl import load_workbook as real_load_workbook

    from st_cost_uploader import parser as parser_module

    closed: list[bool] = []

    def tracking_load_workbook(*args, **kwargs):
        workbook = real_load_workbook(*args, **kwargs)
        original_close = workbook.close

        def close():
            closed.append(True)
            original_close()

        workbook.close = close
        return workbook

    monkeypatch.setattr(parser_module, "load_workbook", tracking_load_workbook)

    result = parse_workbook(
        _xlsx([["Campaign", "Month", "Spend"], ["Yelp", "2026-02", "5000.00"]]),
        "spend.xlsx",
    )

    assert closed == [True]
    # Proof the rows survived the close rather than coming back empty.
    assert [r.campaign_name for r in result.rows] == ["Yelp"]
    assert result.rows[0].monthly_total == Decimal("5000.00")
